#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
VU Exam Exporter CLI Tool
Automates authentication, subject/exam discovery, student attempt fetching,
item analysis (difficulty p, discrimination r), and comprehensive 5-sheet Excel generation.
"""

import sys
import os
import argparse
import json
import re
import math
import getpass
from datetime import datetime
import urllib3
import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

urllib3.disable_warnings()

BASE_URL = "https://i.vu.ac.th/exam-manage-summative"

def load_env_file():
    """Attempt to load .env from current directory or script directory if present."""
    paths_to_check = [
        os.path.join(os.getcwd(), ".env"),
        os.path.join(os.path.dirname(__file__), ".env"),
        os.path.join(os.path.dirname(os.path.dirname(__file__)), ".env")
    ]
    for env_path in paths_to_check:
        if os.path.isfile(env_path):
            try:
                with open(env_path, "r", encoding="utf-8") as f:
                    for line in f:
                        line = line.strip()
                        if line and not line.startswith("#") and "=" in line:
                            k, v = line.split("=", 1)
                            k = k.strip()
                            v = v.strip().strip("'\"")
                            if k not in os.environ:
                                os.environ[k] = v
            except Exception:
                pass
            break

load_env_file()

def clean_html(raw_html):
    if not raw_html:
        return ""
    soup = BeautifulSoup(str(raw_html), "html.parser")
    for br in soup.find_all("br"):
        br.replace_with("\n")
    for p in soup.find_all("p"):
        p.insert_after("\n")
    text = soup.get_text()
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n\s*\n', '\n', text)
    return text.strip()

def sanitize_filename(name):
    clean = re.sub(r'[\\/*?:"<>|]', '_', name)
    return clean.strip()

class VUClient:
    def __init__(self, username=None, password=None):
        self.username = username or os.getenv("VU_EXAM_USERNAME")
        self.password = password or os.getenv("VU_EXAM_PASSWORD")
        self.session = requests.Session()
        self.logged_in_user = None

    def ensure_credentials(self):
        if not self.username:
            try:
                self.username = input("VU Exam Username: ").strip()
            except (EOFError, KeyboardInterrupt):
                pass
        if not self.password:
            try:
                self.password = getpass.getpass("VU Exam Password: ").strip()
            except (EOFError, KeyboardInterrupt):
                pass

    def login(self):
        self.ensure_credentials()
        if not self.username or not self.password:
            return False, "ไม่พบ Username หรือ Password กรุณาระบุผ่าน --username / --password หรือผ่านไฟล์ .env"

        login_url = f"{BASE_URL}/service/select_login.php"
        payload = {"username": self.username, "password": self.password}
        res = self.session.post(login_url, data=payload)
        try:
            data = res.json()
            if int(data.get("success", 0)) == 1:
                self.logged_in_user = data
                return True, data
            return False, data.get("message", "เข้าสู่ระบบไม่สำเร็จ ตรวจสอบ Username/Password")
        except Exception as e:
            return False, f"Login error: {e}"

    def get_subjects(self):
        url = f"{BASE_URL}/service/get_my_subject.php"
        res = self.session.get(url)
        return res.json()

    def get_exams(self, sem_id, subject_id):
        url = f"{BASE_URL}/service/select_exam.php"
        params = {"semid": sem_id, "subjectid": subject_id}
        res = self.session.get(url, params=params)
        return res.json()

    def get_attempts(self, exam_id):
        url = f"{BASE_URL}/service/select_exam_attempt.php"
        params = {"exam_id": exam_id}
        res = self.session.get(url, params=params)
        return res.json()

    def get_attempt_answers(self, exam_id, preload_attempt_id=None, analysis_only=False):
        url = f"{BASE_URL}/service/select_exam_attempt_answers.php"
        params = {"exam_id": exam_id}
        if analysis_only:
            params["analysis_only"] = 1
        elif preload_attempt_id:
            params["preload_attempt_id"] = preload_attempt_id
        res = self.session.get(url, params=params)
        return res.json()

def generate_excel(exam_meta, students_meta, details_dict, output_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    font_family = "Sarabun"
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    sub_header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    accent_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    
    correct_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    wrong_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    
    white_bold_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    title_font = Font(name=font_family, size=16, bold=True, color="1F4E78")
    section_font = Font(name=font_family, size=13, bold=True, color="1F4E78")
    bold_font = Font(name=font_family, size=10, bold=True, color="000000")
    regular_font = Font(name=font_family, size=10, color="000000")
    correct_font = Font(name=font_family, size=10, bold=True, color="276A3C")
    wrong_font = Font(name=font_family, size=10, bold=True, color="C00000")

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    if not students_meta:
        ws = wb.create_sheet(title="ไม่มีผู้เข้าสอบ")
        ws["A1"] = "ไม่มีข้อมูลผู้เข้าสอบสำหรับการสอบนี้"
        wb.save(output_path)
        return

    first_att_id = students_meta[0]['preload_attempt_id']
    first_det = details_dict.get(first_att_id, {})
    
    categories = first_det.get('categories', [])
    if not categories:
        ws = wb.create_sheet(title="ไม่มีข้อสอบ")
        ws["A1"] = "ไม่มีข้อมูลหมวดหมู่ข้อสอบ"
        wb.save(output_path)
        return

    questions_info = []
    q_counter = 1
    for sec in categories[0].get('sections', []):
        sec_title = clean_html(sec.get('section_title', ''))
        for q in sec.get('questions', []):
            q_id = q.get('question_id')
            q_text = clean_html(q.get('question_text', ''))
            score_max = float(q.get('question_score') or 1.0)
            
            options = []
            correct_opts = []
            for opt in q.get('options', []):
                opt_txt = clean_html(opt.get('option_text', ''))
                is_c = (opt.get('is_correct') == '1')
                desc = clean_html(opt.get('answer_option_description', ''))
                options.append({
                    'id': opt.get('exam_option_id'),
                    'text': opt_txt,
                    'is_correct': is_c,
                    'description': desc
                })
                if is_c:
                    correct_opts.append(opt_txt + (f" ({desc})" if desc else ""))

            p_val = q.get('difficulty_index_p')
            r_val = q.get('discrimination_index_r')

            questions_info.append({
                'q_no': q_counter,
                'section': sec_title,
                'question_id': q_id,
                'question_text': q_text,
                'score_max': score_max,
                'options': options,
                'correct_answer': "\n---\n".join(correct_opts) if correct_opts else "N/A",
                'p_value': float(p_val) if p_val is not None else None,
                'diff_level': q.get('difficulty_level', ''),
                'r_value': float(r_val) if r_val is not None else None,
                'disc_level': q.get('discrimination_level', ''),
                'quality_status': q.get('quality_status', ''),
                'quality_label': q.get('quality_label', ''),
                'quality_recommendation': q.get('quality_recommendation', ''),
                'correct_count': int(q.get('correct_count') or 0),
                'incorrect_count': int(q.get('incorrect_count') or 0)
            })
            q_counter += 1

    student_answers = {}
    for stu in students_meta:
        att_id = stu['preload_attempt_id']
        att_det = details_dict.get(att_id, {})
        student_answers[att_id] = {}
        for sec in att_det.get('categories', [{}])[0].get('sections', []):
            for q in sec.get('questions', []):
                qid = q.get('question_id')
                u_ans = q.get('user_answer', [])
                if u_ans and len(u_ans) > 0:
                    ans0 = u_ans[0]
                    sel_text = clean_html(ans0.get('selected_option_text', ''))
                    is_corr = (str(ans0.get('is_correct')) == '1')
                    ans_score = float(ans0.get('answer_score') or 0.0)
                else:
                    sel_text = "(ไม่ได้ตอบ)"
                    is_corr = False
                    ans_score = 0.0
                student_answers[att_id][qid] = {
                    'selected_text': sel_text,
                    'is_correct': is_corr,
                    'score': ans_score
                }

    # 1. Sheet 1: สรุปภาพรวม
    ws1 = wb.create_sheet(title="สรุปภาพรวม")
    ws1.views.sheetView[0].showGridLines = True

    ws1["A1"] = f"รายงานผลการสอบและคะแนนรายบุคคล (Summative Assessment)"
    ws1["A1"].font = title_font
    sub_title = f"วิชา: {exam_meta.get('subject_name', '-')} | ภาคเรียน: {exam_meta.get('sem_label', '-')}"
    ws1["A2"] = sub_title
    ws1["A2"].font = section_font
    ws1["A3"] = f"การสอบ: {exam_meta.get('exam_name', '-')} (Exam ID: {exam_meta.get('exam_id')}) | คะแนนเต็ม: {exam_meta.get('total_score', 0)} คะแนน"
    ws1["A3"].font = Font(name=font_family, size=11, bold=True, color="595959")

    ws1["A5"] = "รายชื่อผู้เข้าสอบและผลคะแนน"
    ws1["A5"].font = section_font
    
    headers_stu = [
        "ลำดับ", "รหัสนักศึกษา", "ชื่อ - นามสกุล", 
        f"คะแนนที่ได้ ({exam_meta.get('total_score', 0)})", "คิดเป็นร้อยละ (%)", 
        "ตอบถูก (ข้อ)", "ตอบผิด (ข้อ)", 
        "เวลาเริ่มสอบ", "เวลาส่งข้อสอบ", "เวลาที่ใช้ (นาที)", "สถานะ"
    ]
    
    for col_idx, h in enumerate(headers_stu, 1):
        cell = ws1.cell(row=6, column=col_idx, value=h)
        cell.font = white_bold_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws1.row_dimensions[6].height = 28

    scores = []
    max_exam_score = float(exam_meta.get('total_score') or (len(questions_info) if questions_info else 100))

    for row_idx, stu in enumerate(students_meta, 7):
        att_id = stu['preload_attempt_id']
        sc = stu['score']
        scores.append(sc)
        pct = (sc / max_exam_score) * 100 if max_exam_score > 0 else 0
        
        c_cnt = sum(1 for q in questions_info if student_answers[att_id].get(q['question_id'], {}).get('is_correct'))
        w_cnt = len(questions_info) - c_cnt

        dur_str = "-"
        if stu['start_time'] and stu['end_time']:
            try:
                t1 = datetime.strptime(stu['start_time'][:19], "%Y-%m-%d %H:%M:%S")
                t2 = datetime.strptime(stu['end_time'][:19], "%Y-%m-%d %H:%M:%S")
                dur_str = f"{round((t2 - t1).total_seconds() / 60, 1):.1f}"
            except Exception:
                pass

        row_vals = [
            row_idx - 6,
            stu['user_id'],
            stu['full_name'],
            sc,
            f"{pct:.2f}%",
            c_cnt,
            w_cnt,
            stu['start_time'],
            stu['end_time'],
            dur_str,
            "ส่งข้อสอบแล้ว (Submitted)" if stu['status'] == "submitted" else stu['status']
        ]

        fill_color = zebra_fill if row_idx % 2 == 0 else PatternFill(fill_type=None)
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if fill_color.fill_type:
                cell.fill = fill_color
            if col_idx in [1, 2, 4, 5, 6, 7, 8, 9, 10, 11]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            
            if col_idx == 4:
                cell.font = bold_font
            elif col_idx == 6:
                cell.font = correct_font
            elif col_idx == 7:
                cell.font = wrong_font
        ws1.row_dimensions[row_idx].height = 22

    avg_score = sum(scores) / len(scores) if scores else 0
    max_sc = max(scores) if scores else 0
    min_sc = min(scores) if scores else 0
    variance = sum((x - avg_score) ** 2 for x in scores) / len(scores) if scores else 0
    sd_score = math.sqrt(variance)

    stat_start_row = len(students_meta) + 9
    ws1.cell(row=stat_start_row, column=1, value="สถิติภาพรวมของการสอบ").font = section_font
    
    stat_items = [
        ("จำนวนผู้เข้าสอบทั้งหมด", f"{len(scores)} คน"),
        ("คะแนนเต็ม", f"{max_exam_score:.0f} คะแนน"),
        ("คะแนนสูงสุด (Max Score)", f"{max_sc:.2f} คะแนน ({(max_sc/max_exam_score)*100:.2f}%)"),
        ("คะแนนต่ำสุด (Min Score)", f"{min_sc:.2f} คะแนน ({(min_sc/max_exam_score)*100:.2f}%)"),
        ("คะแนนเฉลี่ย (Mean Score)", f"{avg_score:.2f} คะแนน ({(avg_score/max_exam_score)*100:.2f}%)"),
        ("ส่วนเบี่ยงเบนมาตรฐาน (S.D.)", f"{sd_score:.2f}")
    ]

    for s_idx, (label, s_val) in enumerate(stat_items, stat_start_row + 1):
        c1 = ws1.cell(row=s_idx, column=1, value=label)
        c2 = ws1.cell(row=s_idx, column=2, value=s_val)
        c1.font = bold_font
        c2.font = regular_font
        c1.fill = accent_fill
        c1.border = thin_border
        c2.border = thin_border
        c1.alignment = Alignment(horizontal="left", vertical="center")
        c2.alignment = Alignment(horizontal="left", vertical="center")
        ws1.row_dimensions[s_idx].height = 20

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)
    ws1.column_dimensions['A'].width = 24
    ws1.column_dimensions['C'].width = 28
    ws1.column_dimensions['H'].width = 22
    ws1.column_dimensions['I'].width = 22

    # 2. Sheet 2: วิเคราะห์รายข้อ (p, r)
    ws2 = wb.create_sheet(title="วิเคราะห์รายข้อ (p, r)")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = "ตารางวิเคราะห์คุณภาพข้อสอบรายข้อ (Item Analysis: Difficulty & Discrimination)"
    ws2["A1"].font = title_font
    ws2["A2"] = "วิเคราะห์ค่าความยากง่าย (p-value) และอำนาจจำแนก (Point-biserial r) พร้อมเกณฑ์และคำแนะนำ"
    ws2["A2"].font = Font(name=font_family, size=11, color="595959")

    headers_item = [
        "ข้อที่", "หัวข้อ (Section)", "รหัสข้อสอบ (ID)", "โจทย์คำถาม", 
        "เฉลยคำตอบที่ถูกต้อง", "คะแนนเต็ม", 
        "ตอบถูก (คน)", "ตอบผิด (คน)", 
        "ค่าความยาก (p)", "ระดับความยาก", 
        "ค่าอำนาจจำแนก (r)", "ระดับอำนาจจำแนก", 
        "สถานะคุณภาพ", "คำแนะนำคุณภาพ"
    ]

    for col_idx, h in enumerate(headers_item, 1):
        cell = ws2.cell(row=4, column=col_idx, value=h)
        cell.font = white_bold_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws2.row_dimensions[4].height = 28

    for row_idx, q in enumerate(questions_info, 5):
        p_str = f"{q['p_value']:.2f}" if q['p_value'] is not None else "-"
        r_str = f"{q['r_value']:.2f}" if q['r_value'] is not None else "-"
        
        row_vals = [
            q['q_no'],
            q['section'],
            q['question_id'],
            q['question_text'],
            q['correct_answer'],
            q['score_max'],
            q['correct_count'],
            q['incorrect_count'],
            p_str,
            q['diff_level'],
            r_str,
            q['disc_level'],
            q['quality_label'] or q['quality_status'],
            q['quality_recommendation']
        ]

        fill_color = zebra_fill if row_idx % 2 == 0 else PatternFill(fill_type=None)
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if fill_color.fill_type:
                cell.fill = fill_color
            
            if col_idx in [1, 3, 6, 7, 8, 9, 11]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_idx in [2, 10, 12, 13]:
                cell.alignment = Alignment(horizontal="left", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            
            if col_idx == 7:
                cell.font = correct_font
            elif col_idx == 8:
                cell.font = wrong_font
            elif col_idx == 13:
                if "ดี" in str(val) or "แนะนำ" in str(val):
                    cell.font = correct_font
                elif "ปรับปรุง" in str(val) or "ไม่ดี" in str(val):
                    cell.font = wrong_font
        ws2.row_dimensions[row_idx].height = 45

    ws2.column_dimensions['A'].width = 8
    ws2.column_dimensions['B'].width = 25
    ws2.column_dimensions['C'].width = 12
    ws2.column_dimensions['D'].width = 45
    ws2.column_dimensions['E'].width = 35
    ws2.column_dimensions['F'].width = 10
    ws2.column_dimensions['G'].width = 12
    ws2.column_dimensions['H'].width = 12
    ws2.column_dimensions['I'].width = 14
    ws2.column_dimensions['J'].width = 20
    ws2.column_dimensions['K'].width = 15
    ws2.column_dimensions['L'].width = 18
    ws2.column_dimensions['M'].width = 16
    ws2.column_dimensions['N'].width = 40

    # 3. Sheet 3: เมทริกซ์คะแนน (1 หรือ 0)
    ws3 = wb.create_sheet(title="เมทริกซ์คะแนน (1 หรือ 0)")
    ws3.views.sheetView[0].showGridLines = True

    ws3["A1"] = "ตารางเมทริกซ์ผลคะแนนรายข้อ (Score Matrix: 1 = ถูก, 0 = ผิด)"
    ws3["A1"].font = title_font

    ws3.cell(row=3, column=1, value="ลำดับ").font = white_bold_font
    ws3.cell(row=3, column=1).fill = header_fill
    ws3.cell(row=3, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws3.cell(row=3, column=1).border = thin_border

    ws3.cell(row=3, column=2, value="รหัสนักศึกษา").font = white_bold_font
    ws3.cell(row=3, column=2).fill = header_fill
    ws3.cell(row=3, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws3.cell(row=3, column=2).border = thin_border

    ws3.cell(row=3, column=3, value="ชื่อ - นามสกุล").font = white_bold_font
    ws3.cell(row=3, column=3).fill = header_fill
    ws3.cell(row=3, column=3).alignment = Alignment(horizontal="center", vertical="center")
    ws3.cell(row=3, column=3).border = thin_border

    for q_idx, q in enumerate(questions_info, 4):
        c = ws3.cell(row=3, column=q_idx, value=f"ข้อ {q['q_no']}")
        c.font = white_bold_font
        c.fill = sub_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        col_letter = get_column_letter(q_idx)
        ws3.column_dimensions[col_letter].width = 7

    total_col_idx = 4 + len(questions_info)
    c_tot = ws3.cell(row=3, column=total_col_idx, value="รวมคะแนน")
    c_tot.font = white_bold_font
    c_tot.fill = header_fill
    c_tot.alignment = Alignment(horizontal="center", vertical="center")
    c_tot.border = thin_border
    ws3.column_dimensions[get_column_letter(total_col_idx)].width = 12
    ws3.row_dimensions[3].height = 26

    for s_idx, stu in enumerate(students_meta, 4):
        att_id = stu['preload_attempt_id']
        ws3.cell(row=s_idx, column=1, value=s_idx - 3).alignment = Alignment(horizontal="center", vertical="center")
        ws3.cell(row=s_idx, column=1).border = thin_border
        ws3.cell(row=s_idx, column=2, value=stu['user_id']).alignment = Alignment(horizontal="center", vertical="center")
        ws3.cell(row=s_idx, column=2).border = thin_border
        ws3.cell(row=s_idx, column=3, value=stu['full_name']).alignment = Alignment(horizontal="left", vertical="center")
        ws3.cell(row=s_idx, column=3).border = thin_border

        tot_s = 0
        for q_idx, q in enumerate(questions_info, 4):
            ans = student_answers[att_id].get(q['question_id'], {})
            is_c = ans.get('is_correct', False)
            sc = 1 if is_c else 0
            tot_s += sc
            
            c = ws3.cell(row=s_idx, column=q_idx, value=sc)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border
            if is_c:
                c.fill = correct_fill
                c.font = correct_font
            else:
                c.fill = wrong_fill
                c.font = wrong_font
        
        c_tot_stu = ws3.cell(row=s_idx, column=total_col_idx, value=tot_s)
        c_tot_stu.alignment = Alignment(horizontal="center", vertical="center")
        c_tot_stu.font = Font(name=font_family, size=11, bold=True, color="1F4E78")
        c_tot_stu.border = thin_border
        ws3.row_dimensions[s_idx].height = 20

    sum_row_idx = 4 + len(students_meta)
    ws3.cell(row=sum_row_idx, column=1, value="").border = thin_border
    ws3.cell(row=sum_row_idx, column=2, value="สรุป").font = bold_font
    ws3.cell(row=sum_row_idx, column=2).border = thin_border
    ws3.cell(row=sum_row_idx, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws3.cell(row=sum_row_idx, column=3, value=f"จำนวนคนตอบถูก (จาก {len(students_meta)})").font = bold_font
    ws3.cell(row=sum_row_idx, column=3).border = thin_border

    for q_idx, q in enumerate(questions_info, 4):
        c = ws3.cell(row=sum_row_idx, column=q_idx, value=q['correct_count'])
        c.font = bold_font
        c.fill = accent_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    
    ws3.cell(row=sum_row_idx, column=total_col_idx, value=f"เฉลี่ย {avg_score:.1f}").font = bold_font
    ws3.cell(row=sum_row_idx, column=total_col_idx).fill = accent_fill
    ws3.cell(row=sum_row_idx, column=total_col_idx).border = thin_border
    ws3.row_dimensions[sum_row_idx].height = 22

    pct_row_idx = sum_row_idx + 1
    ws3.cell(row=pct_row_idx, column=3, value="ร้อยละความถูกต้อง (%)").font = bold_font
    ws3.cell(row=pct_row_idx, column=3).border = thin_border
    for q_idx, q in enumerate(questions_info, 4):
        pct_q = (q['correct_count'] / len(students_meta)) * 100 if students_meta else 0
        c = ws3.cell(row=pct_row_idx, column=q_idx, value=f"{pct_q:.0f}%")
        c.font = regular_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws3.row_dimensions[pct_row_idx].height = 20

    ws3.column_dimensions['A'].width = 8
    ws3.column_dimensions['B'].width = 16
    ws3.column_dimensions['C'].width = 25

    # 4. Sheet 4: เมทริกซ์ตัวเลือกที่ตอบ
    ws4 = wb.create_sheet(title="เมทริกซ์ตัวเลือกที่ตอบ")
    ws4.views.sheetView[0].showGridLines = True

    ws4["A1"] = "ตารางสรุปตัวเลือกคำตอบที่นักศึกษาแต่ละคนเลือก (Selected Options Matrix)"
    ws4["A1"].font = title_font
    ws4["A2"] = "แสดงข้อความตัวเลือกที่เลือกตอบจริง พร้อมสถานะ [✓ ถูก] หรือ [✗ ผิด]"
    ws4["A2"].font = Font(name=font_family, size=11, color="595959")

    ws4.cell(row=4, column=1, value="ลำดับ").font = white_bold_font
    ws4.cell(row=4, column=1).fill = header_fill
    ws4.cell(row=4, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws4.cell(row=4, column=1).border = thin_border

    ws4.cell(row=4, column=2, value="รหัสนักศึกษา").font = white_bold_font
    ws4.cell(row=4, column=2).fill = header_fill
    ws4.cell(row=4, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws4.cell(row=4, column=2).border = thin_border

    ws4.cell(row=4, column=3, value="ชื่อ - นามสกุล").font = white_bold_font
    ws4.cell(row=4, column=3).fill = header_fill
    ws4.cell(row=4, column=3).alignment = Alignment(horizontal="center", vertical="center")
    ws4.cell(row=4, column=3).border = thin_border

    for q_idx, q in enumerate(questions_info, 4):
        c = ws4.cell(row=4, column=q_idx, value=f"ข้อ {q['q_no']}")
        c.font = white_bold_font
        c.fill = sub_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        col_letter = get_column_letter(q_idx)
        ws4.column_dimensions[col_letter].width = 25
    ws4.row_dimensions[4].height = 26

    for s_idx, stu in enumerate(students_meta, 5):
        att_id = stu['preload_attempt_id']
        ws4.cell(row=s_idx, column=1, value=s_idx - 4).alignment = Alignment(horizontal="center", vertical="center")
        ws4.cell(row=s_idx, column=1).border = thin_border
        ws4.cell(row=s_idx, column=2, value=stu['user_id']).alignment = Alignment(horizontal="center", vertical="center")
        ws4.cell(row=s_idx, column=2).border = thin_border
        ws4.cell(row=s_idx, column=3, value=stu['full_name']).alignment = Alignment(horizontal="left", vertical="center")
        ws4.cell(row=s_idx, column=3).border = thin_border

        for q_idx, q in enumerate(questions_info, 4):
            ans = student_answers[att_id].get(q['question_id'], {})
            is_c = ans.get('is_correct', False)
            txt = ans.get('selected_text', '-')
            
            display_val = f"{'✓ ' if is_c else '✗ '}{txt}"
            c = ws4.cell(row=s_idx, column=q_idx, value=display_val)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c.border = thin_border
            if is_c:
                c.fill = correct_fill
                c.font = correct_font
            else:
                c.fill = wrong_fill
                c.font = wrong_font
        ws4.row_dimensions[s_idx].height = 36

    ws4.column_dimensions['A'].width = 8
    ws4.column_dimensions['B'].width = 16
    ws4.column_dimensions['C'].width = 25

    # 5. Sheet 5: รายละเอียดข้อสอบและคำตอบ
    ws5 = wb.create_sheet(title="รายละเอียดข้อสอบและคำตอบ")
    ws5.views.sheetView[0].showGridLines = True

    ws5["A1"] = "รายละเอียดข้อสอบ ตัวเลือก เฉลย และคำตอบของนักศึกษาทุกคนแบบเทียบเคียง"
    ws5["A1"].font = title_font

    headers_breakdown = [
        "ข้อที่", "หมวดหมู่ / Section", "โจทย์คำถาม", 
        "ตัวเลือกทั้งหมด", "เฉลยคำตอบที่ถูกต้อง", "ความยาก (p)", "อำนาจจำแนก (r)"
    ]
    for stu in students_meta:
        headers_breakdown.append(f"{stu['full_name']}\n({stu['user_id']})")

    for col_idx, h in enumerate(headers_breakdown, 1):
        cell = ws5.cell(row=3, column=col_idx, value=h)
        cell.font = white_bold_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws5.row_dimensions[3].height = 32

    for row_idx, q in enumerate(questions_info, 4):
        opts_formatted = []
        for o_i, o in enumerate(q['options']):
            tag = "[✓]" if o['is_correct'] else "[ ]"
            opts_formatted.append(f"{tag} {o['text']}")
        opts_str = "\n".join(opts_formatted)

        p_str = f"{q['p_value']:.2f}" if q['p_value'] is not None else "-"
        r_str = f"{q['r_value']:.2f}" if q['r_value'] is not None else "-"

        row_vals = [
            q['q_no'],
            q['section'],
            q['question_text'],
            opts_str,
            q['correct_answer'],
            p_str,
            r_str
        ]

        for stu in students_meta:
            att_id = stu['preload_attempt_id']
            ans = student_answers[att_id].get(q['question_id'], {})
            is_c = ans.get('is_correct', False)
            t = ans.get('selected_text', '-')
            res_str = f"{'✓ ถูก (1.0)' if is_c else '✗ ผิด (0.0)'}\nตอบ: {t}"
            row_vals.append(res_str)

        fill_color = zebra_fill if row_idx % 2 == 0 else PatternFill(fill_type=None)
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws5.cell(row=row_idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if fill_color.fill_type:
                cell.fill = fill_color
            
            if col_idx in [1, 6, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_idx in [2]:
                cell.alignment = Alignment(horizontal="left", vertical="top")
            elif col_idx in [3, 4, 5]:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                if "✓ ถูก" in str(val):
                    cell.font = correct_font
                    cell.fill = correct_fill
                else:
                    cell.font = wrong_font
                    cell.fill = wrong_fill
        ws5.row_dimensions[row_idx].height = 55

    ws5.column_dimensions['A'].width = 8
    ws5.column_dimensions['B'].width = 24
    ws5.column_dimensions['C'].width = 45
    ws5.column_dimensions['D'].width = 40
    ws5.column_dimensions['E'].width = 35
    ws5.column_dimensions['F'].width = 12
    ws5.column_dimensions['G'].width = 14
    for s_i in range(8, 8 + len(students_meta)):
        ws5.column_dimensions[get_column_letter(s_i)].width = 28

    wb.save(output_path)
    print(f"Excel report saved successfully to: {output_path}")

def cmd_list_subjects(client, args):
    success, res = client.login()
    if not success:
        print(f"Error: {res}")
        sys.exit(1)
    
    subjects = client.get_subjects()
    if args.json:
        print(json.dumps(subjects, ensure_ascii=False, indent=2))
        return

    print("\n" + "="*70)
    print("📚 รายวิชาและภาคเรียนที่สามารถเข้าถึงได้ในระบบ VU Exam Manage Summative")
    print("="*70)
    for idx, sub in enumerate(subjects, 1):
        sem_id = sub.get("open_sem_id")
        term = sub.get("open_sem_term")
        year = sub.get("open_sem_year")
        sub_id = sub.get("subject_id")
        code = sub.get("subject_code")
        name = sub.get("subject_name", "").strip()
        print(f"{idx:2d}. [ภาคเรียน {term}/{year} | semid: {sem_id}] {code} - {name} (subject_id: {sub_id})")
    print("="*70)

def cmd_list_exams(client, args):
    success, res = client.login()
    if not success:
        print(f"Error: {res}")
        sys.exit(1)

    exams = client.get_exams(args.semid, args.subjectid)
    if args.json:
        print(json.dumps(exams, ensure_ascii=False, indent=2))
        return

    print("\n" + "="*70)
    print(f"📝 รายการการสอบสำหรับภาคเรียน {args.semid} | วิชา {args.subjectid}")
    print("="*70)
    if not exams:
        print("  (ไม่พบชุดการสอบในรายวิชานี้)")
    for idx, ex in enumerate(exams, 1):
        eid = ex.get("exam_id")
        ename = ex.get("exam_name")
        status = ex.get("status")
        tot = ex.get("total_score")
        mode = ex.get("exam_mode_name")
        print(f"{idx:2d}. [Exam ID: {eid}] {ename}")
        print(f"    - คะแนนเต็ม: {tot} | สถานะ: {status} | รูปแบบ: {mode}")
    print("="*70)

def cmd_export(client, args):
    success, res = client.login()
    if not success:
        print(f"Error: {res}")
        sys.exit(1)

    exam_id = args.exam_id
    print(f"[*] กำลังดึงข้อมูลการสอบ Exam ID: {exam_id}...")
    
    r_att = client.get_attempts(exam_id)
    stu_raw = r_att.get("data", [])
    print(f"[*] พบผู้เข้าสอบทั้งหมด: {len(stu_raw)} คน")

    students_meta = []
    for stu in stu_raw:
        uid = stu.get('user_id')
        name = stu.get('full_name')
        for att in stu.get('attempts', []):
            att_id = att.get('preload_attempt_id')
            score = att.get('total_score')
            max_score = att.get('exam_total_score')
            students_meta.append({
                'user_id': uid,
                'full_name': name,
                'preload_attempt_id': att_id,
                'score': float(score) if score is not None else 0.0,
                'max_score': float(max_score) if max_score is not None else 0.0,
                'status': att.get('status'),
                'start_time': att.get('start_time'),
                'end_time': att.get('end_time')
            })

    details_dict = {}
    exam_info_meta = {'exam_id': exam_id}

    for idx, s in enumerate(students_meta, 1):
        att_id = s['preload_attempt_id']
        name = s['full_name']
        print(f"    [{idx}/{len(students_meta)}] กำลังโหลดผลการสอบของ {name} (Attempt: {att_id})...")
        d = client.get_attempt_answers(exam_id, preload_attempt_id=att_id)
        det_data = d.get('data', {})
        details_dict[att_id] = det_data
        
        if not exam_info_meta.get('exam_name') and det_data.get('exam_name'):
            exam_info_meta['exam_name'] = det_data.get('exam_name')
            exam_info_meta['total_score'] = det_data.get('total_score')

    exam_title_clean = sanitize_filename(exam_info_meta.get('exam_name', f'Exam_{exam_id}'))
    output_filename = args.output or f"VU_Exam_{exam_id}_{exam_title_clean}.xlsx"

    print("[*] กำลังสร้างไฟล์ Excel รายงาน 5 แผ่นงาน...")
    generate_excel(exam_info_meta, students_meta, details_dict, output_filename)
    print(f"[✓] เสร็จสมบูรณ์! บันทึกไฟล์ที่: {os.path.abspath(output_filename)}")

def main():
    parser = argparse.ArgumentParser(description="VU Exam Summative Exporter CLI")
    parser.add_argument("--username", "-u", default=None, help="VU Exam Username (or via VU_EXAM_USERNAME env)")
    parser.add_argument("--password", "-p", default=None, help="VU Exam Password (or via VU_EXAM_PASSWORD env)")

    subparsers = parser.add_subparsers(dest="command", required=True)

    p_subjects = subparsers.add_parser("list-subjects", help="List all available semesters and subjects")
    p_subjects.add_argument("--json", action="store_true", help="Output as raw JSON")

    p_exams = subparsers.add_parser("list-exams", help="List all exams for a specific semester and subject")
    p_exams.add_argument("--semid", required=True, help="Semester ID (e.g. 25691)")
    p_exams.add_argument("--subjectid", required=True, help="Subject ID (e.g. 14)")
    p_exams.add_argument("--json", action="store_true", help="Output as raw JSON")

    p_export = subparsers.add_parser("export", help="Export full exam data and item analysis to Excel")
    p_export.add_argument("--exam-id", required=True, help="Exam ID (e.g. 1034)")
    p_export.add_argument("--output", "-o", help="Output .xlsx file path")

    args = parser.parse_args()
    client = VUClient(args.username, args.password)

    if args.command == "list-subjects":
        cmd_list_subjects(client, args)
    elif args.command == "list-exams":
        cmd_list_exams(client, args)
    elif args.command == "export":
        cmd_export(client, args)

if __name__ == "__main__":
    main()
