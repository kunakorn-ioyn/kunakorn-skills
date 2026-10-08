#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Full System & KM Report DOCX Generator
Generates high-quality, professional Word documents (.docx) for Knowledge Management (KM)
and System Technical Specifications with standard formatting, tables, and callouts.
"""

import os
import sys
import json
import argparse
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

PRIMARY_COLOR = RGBColor(30, 58, 138)    # Deep Blue #1E3A8A
SECONDARY_COLOR = RGBColor(79, 70, 229)  # Indigo #4F46E5
TEXT_DARK = RGBColor(31, 41, 55)         # Charcoal #1F2937
MUTED_TEXT = RGBColor(107, 114, 128)     # Slate Gray #6B7280

FONT_NAME = "TH Sarabun New"
FONT_NAME_ALT = "Angsana New"

def set_cell_background(cell, hex_color):
    """Set shading/background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set internal cell padding (in dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout(doc, text_content, title="ข้อความสำคัญ / สรุปสาระสำคัญ"):
    """Adds a stylish callout box with a shaded background."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border styling (thick blue left border)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="1E3A8A"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(f"📌 {title}\n")
    r_title.bold = True
    r_title.font.name = FONT_NAME
    r_title.font.size = Pt(15)
    r_title.font.color.rgb = PRIMARY_COLOR
    
    r_body = p.add_run(text_content)
    r_body.font.name = FONT_NAME
    r_body.font.size = Pt(14)
    r_body.font.color.rgb = TEXT_DARK
    
    # spacing after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)

def format_run_font(run, font_name=FONT_NAME, size_pt=16, color=TEXT_DARK, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic

def generate_report_from_json(data, output_path):
    """Generates the .docx report from a structured Python dict."""
    doc = docx.Document()
    
    # Page setup: A4 with 1-inch margins
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / Footer
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run(f"เอกสารรายงานระบบ: {data.get('title', 'System Report')} | หน้า ")
        format_run_font(r_ft, size_pt=11, color=MUTED_TEXT)

    # --- Title Page / Hero Header ---
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(18)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run(data.get("title", "รายงานวิศวกรรมระบบและจัดการความรู้"))
    format_run_font(r_title, size_pt=24, color=PRIMARY_COLOR, bold=True)
    
    if data.get("subtitle"):
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(12)
        r_sub = p_sub.add_run(data["subtitle"])
        format_run_font(r_sub, size_pt=16, color=SECONDARY_COLOR, bold=True)
        
    # Metadata Box (Key details)
    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_items = data.get("metadata", {})
    
    for k, v in meta_items.items():
        row = meta_table.add_row()
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "F8FAFC")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(str(k))
        format_run_font(r0, size_pt=13, color=PRIMARY_COLOR, bold=True)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(str(v))
        format_run_font(r1, size_pt=13, color=TEXT_DARK)
        
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(12)
    
    # Executive Summary / Callout
    if data.get("executive_summary"):
        add_callout(doc, data["executive_summary"], title="บทสรุปสำหรับผู้บริหาร (Executive Summary)")

    # --- Chapters & Content ---
    chapters = data.get("chapters", [])
    for idx, chap in enumerate(chapters, start=1):
        chap_title = f"{idx}. {chap.get('title', '')}" if chap.get('numbering', True) else chap.get('title', '')
        
        # Heading 1
        p_h1 = doc.add_paragraph()
        p_h1.paragraph_format.space_before = Pt(16)
        p_h1.paragraph_format.space_after = Pt(6)
        p_h1.paragraph_format.keep_with_next = True
        r_h1 = p_h1.add_run(chap_title)
        format_run_font(r_h1, size_pt=18, color=PRIMARY_COLOR, bold=True)
        
        # Chapter paragraphs
        for p_data in chap.get("content", []):
            if isinstance(p_data, str):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run(p_data)
                format_run_font(r, size_pt=14, color=TEXT_DARK)
            elif isinstance(p_data, dict):
                # Subheading 2
                if p_data.get("type") == "h2":
                    p_h2 = doc.add_paragraph()
                    p_h2.paragraph_format.space_before = Pt(12)
                    p_h2.paragraph_format.space_after = Pt(4)
                    p_h2.paragraph_format.keep_with_next = True
                    r_h2 = p_h2.add_run(p_data.get("text", ""))
                    format_run_font(r_h2, size_pt=16, color=SECONDARY_COLOR, bold=True)
                # Subheading 3
                elif p_data.get("type") == "h3":
                    p_h3 = doc.add_paragraph()
                    p_h3.paragraph_format.space_before = Pt(8)
                    p_h3.paragraph_format.space_after = Pt(2)
                    p_h3.paragraph_format.keep_with_next = True
                    r_h3 = p_h3.add_run(p_data.get("text", ""))
                    format_run_font(r_h3, size_pt=14.5, color=TEXT_DARK, bold=True)
                # Callout
                elif p_data.get("type") == "callout":
                    add_callout(doc, p_data.get("text", ""), title=p_data.get("title", "สาระสำคัญ"))
                # Bullet list
                elif p_data.get("type") == "bullet":
                    p = doc.add_paragraph(style='List Bullet')
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    r = p.add_run(p_data.get("text", ""))
                    format_run_font(r, size_pt=14, color=TEXT_DARK)
                # Numbered list
                elif p_data.get("type") == "number":
                    p = doc.add_paragraph(style='List Number')
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    r = p.add_run(p_data.get("text", ""))
                    format_run_font(r, size_pt=14, color=TEXT_DARK)
                # Table
                elif p_data.get("type") == "table":
                    tbl_headers = p_data.get("headers", [])
                    tbl_rows = p_data.get("rows", [])
                    
                    t = doc.add_table(rows=1, cols=len(tbl_headers))
                    t.alignment = WD_TABLE_ALIGNMENT.CENTER
                    
                    # Headers
                    hdr_cells = t.rows[0].cells
                    for i, h in enumerate(tbl_headers):
                        hdr_cells[i].text = h
                        set_cell_background(hdr_cells[i], "1E3A8A")
                        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=120, right=120)
                        p = hdr_cells[i].paragraphs[0]
                        for run in p.runs:
                            format_run_font(run, size_pt=13, color=RGBColor(255, 255, 255), bold=True)
                            
                    # Rows
                    for r_idx, r_vals in enumerate(tbl_rows):
                        row = t.add_row()
                        bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
                        for c_idx, val in enumerate(r_vals):
                            cell = row.cells[c_idx]
                            cell.text = str(val)
                            set_cell_background(cell, bg_color)
                            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
                            p = cell.paragraphs[0]
                            for run in p.runs:
                                format_run_font(run, size_pt=13, color=TEXT_DARK)
                                
                    p_sp = doc.add_paragraph()
                    p_sp.paragraph_format.space_after = Pt(6)

    # Save
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    print(f"✅ Full Report successfully generated: {output_path}")

def main():
    parser = argparse.ArgumentParser(description="Generate Word (.docx) Full System / KM Report")
    parser.add_argument("--input-json", required=True, help="Path to input JSON report structure")
    parser.add_argument("--output", required=True, help="Output .docx file path")
    args = parser.parse_args()
    
    with open(args.input_json, "r", encoding="utf-8") as f:
        data = json.load(f)
        
    generate_report_from_json(data, args.output)

if __name__ == "__main__":
    main()
