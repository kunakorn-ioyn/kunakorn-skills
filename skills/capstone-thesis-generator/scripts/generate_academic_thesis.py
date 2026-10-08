#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Formal Academic Capstone Thesis DOCX Engine (Thai University Standard)
Strictly adheres to formal Thai university project report standards:
- Zero emojis
- Formal black typography (RGB 0, 0, 0)
- TH Sarabun New font family (16pt body, 16pt/18pt bold headings)
- Standard academic table borders with subtle shaded headers
- Paragraph first-line indentation (0.5 inch / 1.27 cm)
- Standard page margins (Left 1.5 inches for book binding, Top/Right/Bottom 1.0 inch)
- Proper Chapter, Section, Table, and Figure captions
- High-resolution figure image embedding
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BLACK = RGBColor(0, 0, 0)
GRAY_BORDER = "CCCCCC"
GRAY_HEADER_BG = "F3F4F6"
ALT_ROW_BG = "FAFAFA"
FONT_NAME = "TH Sarabun New"

def set_cell_background(cell, hex_color):
    """Set background shading color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    """Set internal cell padding (in dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table):
    """Apply standard formal borders to a table (top/bottom solid black, no vertical borders)."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{GRAY_BORDER}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_run(run, font_name=FONT_NAME, size_pt=16, color=BLACK, bold=False, italic=False):
    """Apply consistent Thai academic typography to a run."""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
    rPr.append(rFonts)

def generate_academic_thesis_docx(data, output_path):
    """Generates the formal academic thesis document."""
    doc = docx.Document()
    
    # Page setup: A4 with standard margins (Left 1.5 inch for binding, Top/Right/Bottom 1.0 inch)
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)
        
        # Header / Footer
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run(f"{data.get('short_title', 'รายงานโครงงานวิจัย')} ")
        format_run(r_ft, size_pt=12, color=RGBColor(120, 120, 120))

    # --- Cover Page ---
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.paragraph_format.line_spacing = 1.15
    r_title = p_title.add_run(data.get("title", ""))
    format_run(r_title, size_pt=20, bold=True)
    
    if data.get("title_en"):
        p_en = doc.add_paragraph()
        p_en.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_en.paragraph_format.space_after = Pt(24)
        r_en = p_en.add_run(data["title_en"])
        format_run(r_en, size_pt=16, bold=True)
        
    p_author_block = doc.add_paragraph()
    p_author_block.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author_block.paragraph_format.space_before = Pt(72)
    p_author_block.paragraph_format.space_after = Pt(6)
    r_author = p_author_block.add_run(data.get("author", "คณะผู้จัดทำโครงงาน"))
    format_run(r_author, size_pt=16, bold=True)

    p_org = doc.add_paragraph()
    p_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_org.paragraph_format.space_after = Pt(4)
    r_org = p_org.add_run(data.get("faculty_univ", "มหาวิทยาลัย"))
    format_run(r_org, size_pt=16)

    p_year = doc.add_paragraph()
    p_year.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_year.paragraph_format.space_after = Pt(72)
    r_year = p_year.add_run(f"ปีการศึกษา {data.get('academic_year', '2569')}")
    format_run(r_year, size_pt=16)

    doc.add_page_break()

    # --- Abstract Thai ---
    if data.get("abstract_th"):
        p_ab_th = doc.add_paragraph()
        p_ab_th.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ab_th.paragraph_format.space_before = Pt(12)
        p_ab_th.paragraph_format.space_after = Pt(12)
        r_ab_th = p_ab_th.add_run("บทคัดย่อ")
        format_run(r_ab_th, size_pt=18, bold=True)

        for p_text in data["abstract_th"]:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(0.5)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(p_text)
            format_run(r, size_pt=16)

        doc.add_page_break()

    # --- Abstract English ---
    if data.get("abstract_en"):
        p_ab_en = doc.add_paragraph()
        p_ab_en.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ab_en.paragraph_format.space_before = Pt(12)
        p_ab_en.paragraph_format.space_after = Pt(12)
        r_ab_en = p_ab_en.add_run("Abstract")
        format_run(r_ab_en, size_pt=18, bold=True)

        for p_text in data["abstract_en"]:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(0.5)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(p_text)
            format_run(r, size_pt=16)

        doc.add_page_break()

    # --- Chapters ---
    chapters = data.get("chapters", [])
    for c_idx, chap in enumerate(chapters, start=1):
        # Chapter Title
        p_c_title = doc.add_paragraph()
        p_c_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_c_title.paragraph_format.space_before = Pt(24)
        p_c_title.paragraph_format.space_after = Pt(6)
        r_c_num = p_c_title.add_run(chap.get("header_title", f"บทที่ {c_idx}\n{chap.get('title', '')}"))
        format_run(r_c_num, size_pt=18, bold=True)

        for item in chap.get("content", []):
            if isinstance(item, str):
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Inches(0.5)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run(item)
                format_run(r, size_pt=16)
            elif isinstance(item, dict):
                i_type = item.get("type")
                # Heading 2
                if i_type == "h2":
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(14)
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.keep_with_next = True
                    r = p.add_run(item.get("text", ""))
                    format_run(r, size_pt=16, bold=True)
                # Heading 3
                elif i_type == "h3":
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(10)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.keep_with_next = True
                    r = p.add_run(item.get("text", ""))
                    format_run(r, size_pt=16, bold=True)
                # Bullet list
                elif i_type == "bullet":
                    p = doc.add_paragraph(style='List Bullet')
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    r = p.add_run(item.get("text", ""))
                    format_run(r, size_pt=16)
                # Numbered list
                elif i_type == "number":
                    p = doc.add_paragraph(style='List Number')
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    r = p.add_run(item.get("text", ""))
                    format_run(r, size_pt=16)
                # Table Caption
                elif i_type == "table_caption":
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(12)
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.keep_with_next = True
                    r = p.add_run(item.get("text", ""))
                    format_run(r, size_pt=15, bold=True)
                # Figure Caption
                elif i_type == "figure_caption":
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(8)
                    r = p.add_run(item.get("text", ""))
                    format_run(r, size_pt=15, bold=True)
                # Figure Image
                elif i_type == "image":
                    img_path = item.get("path")
                    if img_path and os.path.exists(img_path):
                        p_img = doc.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img.paragraph_format.space_before = Pt(8)
                        p_img.paragraph_format.space_after = Pt(2)
                        run_img = p_img.add_run()
                        run_img.add_picture(img_path, width=Inches(item.get("width_inch", 5.5)))
                # Formal Table
                elif i_type == "table":
                    headers = item.get("headers", [])
                    rows = item.get("rows", [])
                    
                    tbl = doc.add_table(rows=1, cols=len(headers))
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(tbl)
                    
                    hdr_cells = tbl.rows[0].cells
                    for col_idx, h_text in enumerate(headers):
                        hdr_cells[col_idx].text = h_text
                        set_cell_background(hdr_cells[col_idx], GRAY_HEADER_BG)
                        set_cell_margins(hdr_cells[col_idx], top=70, bottom=70, left=100, right=100)
                        p_hdr = hdr_cells[col_idx].paragraphs[0]
                        p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in p_hdr.runs:
                            format_run(run, size_pt=14, bold=True)
                            
                    for r_idx, row_vals in enumerate(rows):
                        row = tbl.add_row()
                        row_bg = ALT_ROW_BG if r_idx % 2 == 1 else "FFFFFF"
                        for c_idx, val in enumerate(row_vals):
                            cell = row.cells[c_idx]
                            cell.text = str(val)
                            set_cell_background(cell, row_bg)
                            set_cell_margins(cell, top=60, bottom=60, left=90, right=90)
                            p_cell = cell.paragraphs[0]
                            # Centering for row order / short codes
                            if c_idx == 0 and len(str(val)) <= 4:
                                p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for run in p_cell.runs:
                                format_run(run, size_pt=14)
                                
                    p_sp = doc.add_paragraph()
                    p_sp.paragraph_format.space_after = Pt(6)

        # Page break after each chapter
        doc.add_page_break()

    # Save
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    print(f"Academic Capstone Thesis DOCX successfully generated: {output_path}")

if __name__ == "__main__":
    pass
